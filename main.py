import requests
import time
from dotenv import load_dotenv
import os
import json
from multiprocessing import Pool
import random
import pandas as pd
from openpyxl import load_workbook


load_dotenv()

env_file = ".env.range"
load_dotenv(dotenv_path=env_file)
cookies = json.loads(os.getenv("cookies"))
headers = json.loads(os.getenv("headers"))

url = "https://www.wildberries.ru/__internal/search/exactmatch/ru/common/v18/search?ab_testing=false&appType=1&autoselectFilters=false&curr=rub&dest=-5803648&hide_vflags=4294967296&inheritFilters=false&lang=ru&priceU=27400%3B2000000&query=%D1%82%D0%B5%D0%BB%D0%B5%D1%84%D0%BE%D0%BD%D1%8B&resultset=filters&spp=30&suppressSpellcheck=false&uclusters=2"

def get_total_count(min_p, max_p):
    params = {
        'ab_testing': 'false',
        'appType': '1',
        'autoselectFilters': 'false',
        'curr': 'rub',
        'dest': '-5803648',
        'hide_vflags': '4294967296',
        'inheritFilters': 'false',
        'lang': 'ru',
        'priceU': f"{min_p * 100};{max_p * 100}",
        'query': 'телефоны',
        'resultset': 'filters',
        'spp': '30',
        'suppressSpellcheck': 'false',
        'uclusters': '2',
    }
    try:
        time.sleep(1)
        response = requests.get(url, params=params, headers=headers, cookies=cookies, timeout=5)
        data = response.json()
        return data.get('data', {}).get('total', 0)
    except Exception:
        return 0


def split_range(min_p, max_p):
    total = get_total_count(min_p, max_p)
    print(f"Диапазон {min_p}-{max_p} руб: найдено {total} товаров")

    if total <= 5000:
        line = f"Парсим от {min_p} до {max_p} руб \n"
        with open("range.txt", "a", encoding="utf-8") as f:
            f.write(line)

        return [(min_p, max_p)]
    else:
        middle = (min_p + max_p) // 2

        left_part = split_range(min_p, middle)
        right_part = split_range(middle + 1, max_p)

        return left_part + right_part

def download_json(rau):
    min_u, max_u = rau
    page = 1
    count = 0
    while True:
        try:
            url = f"https://www.wildberries.ru/__internal/search/exactmatch/ru/common/v18/search?ab_testing=false&appType=1&curr=rub&dest=-5803648&hide_vflags=4294967296&inheritFilters=false&lang=ru&priceU={min_u}%3B{max_u}&query=%D1%82%D0%B5%D0%BB%D0%B5%D1%84%D0%BE%D0%BD%D1%8B&resultset=catalog&sort=popular&spp=30&suppressSpellcheck=false&uclusters=2&page={page}"
            params = {
                'ab_testing': 'false',
                'appType': '1',
                'curr': 'rub',
                'dest': '-5803648',
                'hide_vflags': '4294967296',
                'inheritFilters': 'false',
                'lang': 'ru',
                'priceU': f'{min_u};{max_u}',
                'query': 'телефоны',
                'resultset': 'catalog',
                'sort': 'popular',
                'spp': '30',
                'suppressSpellcheck': 'false',
                'uclusters': '2',
            }

            time.sleep(random.uniform(2.0, 4.0))
            req = requests.get(url, cookies=cookies, params=params, headers=headers)
            data = req.json()

            if req.status_code != 200:
                print("Ошибка сервера!")
                break

            if len(data["products"]) == 0:
                print(f"Диапазон {min_u}-{max_u} полностью собран. Страниц: {page - 1}")
                break

            file_name = rf"TelephoneParser\phones{min_u // 100}-{max_u // 100}_{count}.json"
            with open(file_name, "w", encoding="utf-8") as file:
                json.dump(data, file, indent=4, ensure_ascii=False)

            page += 1
            count += 1
        except Exception as e:
            print(f"Произошла ошибка {e}")
            break

if __name__ == '__main__':
    with open("range.txt", "w", encoding="utf-8") as f:
        f.write("")

    final_ranges = split_range(10000, 35000)

    # смотрим диапазоны и по ним достаем json
    price_range_list = []
    with open("range.txt", "r", encoding="utf-8") as f:
        for line in f:
            if "Парсим от" in line:
                parts = line.split()
                min_p = int(parts[2])
                max_p = int(parts[4])
                price_range_list.append([min_p, max_p])

    range_and_url = []
    for i in price_range_list:
        max_u = i[1] * 100
        min_u = i[0] * 100
        range_and_url.append([min_u, max_u])

    with Pool(2) as p:
        p.map(download_json, range_and_url)

    # читаем json которые сохранили
    data = []
    for price_range in price_range_list:
        min_u, max_u = price_range
        count = 0
        while True:
            try:
                file_path = f"TelephoneParser/phones{min_u}-{max_u}_{count}.json"
                with open(file_path, "r", encoding="utf-8") as file:
                    data.append(json.load(file))
            except Exception:
                break
            count += 1

    all_phones = []
    for phones in data:
        for phone in phones.get("products", []):
            meta = phone.get("meta", {})
            chars = meta.get("characteristics", [])
            memory = "Не указано"
            for char in chars:
                if char.get("name") == "Объем встроенной памяти":
                    memory = char.get("values", ["Не указано"])[0]

            all_phones.append({
                "Бренд": phone.get("brand") if phone.get("brand") else "Без бренда",
                "Название": phone.get("name"),
                "Память": memory,
                "Цена": phone["sizes"][0]["price"].get("product", 0) / 100 if phone.get("sizes") else 0,
                "Рейтинг": phone.get("reviewRating"),
                "Отзывов": phone.get("feedbacks"),
                "Продавец": phone.get("supplier"),
                "Рейтинг продавца": phone.get("supplierRating"),
                "Цвет": phone["colors"][0].get("name") if phone.get("colors") else "Нет цвета",
                "Артикул": phone.get("id"),
            })

    # в excel сохраняем
    import pandas as pd
    from openpyxl import load_workbook

    df = pd.DataFrame(all_phones)
    file_name = "Телефоны.xlsx"
    df.to_excel(file_name, index=False)

    # размеры колонок
    wb = load_workbook(file_name)
    ws = wb.active
    ws.column_dimensions['G'].width = 25
    ws.column_dimensions['B'].width = 50
    wb.save(file_name)

    print(f"Готово! Данные сохранены в файл: {file_name}")